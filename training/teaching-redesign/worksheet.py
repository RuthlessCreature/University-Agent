from pathlib import Path
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=Path(__file__).resolve().parent;O=R/'output';O.mkdir(exist_ok=True)
d=Document();sec=d.sections[0];sec.page_height=Cm(29.7);sec.page_width=Cm(21)
sec.top_margin=sec.bottom_margin=Cm(1.8);sec.left_margin=sec.right_margin=Cm(2)
for name in ['Normal','Title','Heading 1','Heading 2']:
    st=d.styles[name];st.font.name='Microsoft YaHei';st._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei');st.font.color.rgb=RGBColor(0,0,0)
d.styles['Normal'].font.size=Pt(10);d.styles['Normal'].paragraph_format.space_after=Pt(5)
d.styles['Normal'].paragraph_format.line_spacing=1.15
d.styles['Title'].font.size=Pt(20)
d.styles['Heading 1'].font.size=Pt(12);d.styles['Heading 1'].paragraph_format.space_before=Pt(9)
d.add_paragraph('AI 辅助文献筛选研究设计工作底稿','Title')
d.add_paragraph('课程讨论用   2026 年 10 月 17 日')
d.add_paragraph('本底稿采用拟议研究和教学算例，含 AI 辅助编写，尚未实施实验或完成查新。它用于检查研究设计，不是基金申请书。学员应以本人研究材料填写并承担学术判断。')
parts=[
('研究问题','初始相关样本的主题覆盖，是否改变主动学习早停时的漏检分布？数据集是否适合检验这一问题、同类研究是否已完成该比较，均待查证。'),
('资料与参考标签','拟使用有许可的公开题名与摘要数据。依据预先确定的纳入标准，由独立筛选者标注并仲裁分歧。记录数据来源、许可、纳入标准、标签版本及缺失情况。当前尚未选定实际数据集。'),
('比较与信息隔离','同一集合比较随机顺序与主动学习排序。预先固定初始标签、模型配置及停止规则，保留随机种子。评价脚本使用完整参考标签；排序引擎只读取已经揭示的标签。调参与最终测试使用不同集合。'),
('评价与解释','拟报告预定召回水平下的筛选比例，以及停止时的漏检记录和主题。课堂暂用 95% 作为评价点，实际阈值须结合漏检后果论证。多次运行保留变异；人工阅读与计算耗时分别记录。'),
('算例与分母','教学设定：总计 1,000 条，相关 100 条；已筛 300 条，找出相关 95 条。召回率 95%，已筛记录中的相关比例约 31.7%，筛选比例 30%，尚漏 5 条。该算例不代表实验结果或已节约的工时。'),
('来源与待补条件','方法背景：van de Schoot 等，Nature Machine Intelligence 3，125–133（2021），DOI 10.1038/s42256-020-00287-7。下一步须补充停止规则和跨主题评价文献、数据许可与标注成本，再决定研究是否可行。'),
]
for title,body in parts:d.add_paragraph(title,'Heading 1');d.add_paragraph(body)
d.add_paragraph('核对记录：研究者、核对日期、资料版本与未解决问题应随真实任务填写。不得填入不存在的研究基础或成果。')
for el in list(d.styles.element.iter(qn('w:pBdr'))):el.getparent().remove(el)
for el in list(d.element.iter(qn('w:pBdr'))):el.getparent().remove(el)
d.save(O/'研究设计工作底稿.docx')
print(O/'研究设计工作底稿.docx')
