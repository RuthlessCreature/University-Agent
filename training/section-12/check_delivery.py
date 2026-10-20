from pathlib import Path
from xml.etree import ElementTree as E
import csv,json,hashlib,zipfile,re,statistics,io,subprocess,sys
from docx import Document
from pypdf import PdfReader
from openpyxl import load_workbook
r=Path(__file__).resolve().parent;rev=sys.argv[1] if len(sys.argv)>1 else 's12-r8';packet=r/'closeout-packet';p=next((r/'output').glob('*'+rev+'.pptx'))
def load(p):return json.loads(p.read_text('utf-8'))
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
run=load(packet/'output/run.json');assert run==load(r/'execution-evidence.json')
assert all(sha(packet/'input'/name)==h for name,h in run['input_hashes'].items())
assert sha(r/'run_closeout.py')==run['executor_sha256'] and sha(packet/'output/materials-review.docx')==run['artifact_sha256']
assert run['status']=='materials_review_required' and not run['human_approved'] and not run['external_model_called'] and not run['official_submission'] and run['scientific_validation']=='not_performed'
with (packet/'input/penguins.csv').open(encoding='utf-8-sig',newline='') as f:raw=list(csv.DictReader(f))
ad=[q for q in raw if q['species']=='Adelie'];clean=[q for q in ad if q['sex'] in ['female','male'] and q['flipper_length_mm']!='NA']
assert [len(raw),len(ad),len(clean)]==[344,152,146]
groups={s:[float(q['flipper_length_mm']) for q in clean if q['sex']==s] for s in ['female','male']};means=[statistics.fmean(groups[s]) for s in ['female','male']]
calc=load(packet/'output/calculation.json');assert all(abs(means[i]-calc[s+'_mean_mm'])<1e-10 for i,s in enumerate(['female','male']))
assert abs(means[1]-means[0]-calc['mean_difference_mm'])<1e-10 and len(groups['female'])==len(groups['male'])==73
matrix=load(packet/'output/evidence-matrix.json');assert [q['status'] for q in matrix]==['VERIFIED','VERIFIED','PENDING']
assert len(load(packet/'output/issues.json'))==2 and all(q['status']=='PENDING' for q in load(packet/'output/issues.json'))
assert len(load(packet/'output/excluded-claims.json'))==2 and all(q['status']=='UNSUPPORTED' for q in load(packet/'output/excluded-claims.json'))
neg=load(r/'negative-cases.json');assert len(neg)==4 and all(q['status']=='failed' and q['word_written']==False for q in neg)
doc=Document(packet/'output/materials-review.docx');body='\n'.join(q.text for q in doc.paragraphs)+'\n'+'\n'.join(c.text for t in doc.tables for row in t.rows for c in row.cells)
assert all(t in body for t in ['M1','M2','M3','187.79','192.41','4.62','AI编写','不能作为该类正式报告提交']) and '{{' not in body
assert len(PdfReader(packet/'word-preview/materials-review.pdf').pages)==3 and len(list((packet/'word-preview').glob('page-*.png')))==3
catalog=load(packet/'output/attachment-catalog.json')['attachments']
assert len(catalog)==run['attachment_count']==9
assert sha(packet/'output/attachment-catalog.json')==run['attachment_catalog_sha256']
assert [a['id'] for a in catalog]==['A'+str(i).zfill(2) for i in range(1,10)]
for a in catalog:
 f=(packet/a['path']).resolve();assert f.is_relative_to(r)
 assert f.exists() and f.stat().st_size==a['bytes'] and sha(f)==a['sha256']
with (packet/'output/attachment-catalog.csv').open(encoding='utf-8-sig',newline='') as f:
 saved=list(csv.DictReader(f))
assert len(saved)==9
for a,b in zip(catalog,saved):assert all(str(a[k])==v for k,v in b.items())
with (packet/'output/retained-rows.csv').open(encoding='utf-8-sig',newline='') as f:kept=list(csv.DictReader(f))
expected=[(i,q) for i,q in enumerate(raw,2) if q['species']=='Adelie' and q['sex'] in ['female','male'] and q['flipper_length_mm'] not in ['', 'NA']]
assert len(kept)==146 and len(set(q['csv_physical_line'] for q in kept))==146
for a,(i,b) in zip(kept,expected):
 assert int(a['csv_physical_line'])==i and a['sex']==b['sex'] and a['flipper_length_mm']==b['flipper_length_mm']
assert all(a['id'] in body for a in catalog) and '教学材料附件目录' in body
lesson=load(r/'lesson.json');ns={'a':'http://schemas.openxmlformats.org/drawingml/2006/main','p':'http://schemas.openxmlformats.org/presentationml/2006/main','c':'http://schemas.openxmlformats.org/drawingml/2006/chart'};diagrams=[];tables=[];charts=[];images=[];texts=[]
with zipfile.ZipFile(p) as z:
 slides=[n for n in z.namelist() if re.fullmatch(r'ppt/slides/slide\d+\.xml',n)];notes=[n for n in z.namelist() if re.fullmatch(r'ppt/notesSlides/notesSlide\d+\.xml',n)];assert len(slides)==len(notes)==len(lesson)==10
 for i in range(1,11):
  x=E.fromstring(z.read(f'ppt/slides/slide{i}.xml'));text=' '.join(t.text or '' for t in x.findall('.//a:t',ns));texts.append(text);assert lesson[i-1]['title'] in text
  assert all(t not in text for t in ['TODO','课堂练习','占位'])
  if len(x.findall('.//p:cxnSp',ns))>=2:diagrams.append(i)
  if x.findall('.//a:tbl',ns):tables.append(i)
  if x.findall('.//c:chart',ns):charts.append(i)
  if x.findall('.//p:pic',ns):images.append(i)
  nt=z.read(f'ppt/notesSlides/notesSlide{i}.xml').decode();assert 'http' in nt or 'local:' in nt
 assert tables==[5,8] and charts==[6] and images==[9]
 flow=E.fromstring(z.read('ppt/slides/slide7.xml'));nodes={}
 for sp in flow.findall('.//p:sp',ns):
  label=''.join(t.text or '' for t in sp.findall('.//a:t',ns))
  nv=sp.find('./p:nvSpPr/p:cNvPr',ns)
  if nv is not None:nodes[label]=nv.attrib['id']
 expected_edges=[('核对来源与文件版本','复算筛选与分组均值'),('复算筛选与分组均值','继续查个体与采样依据'),('复算筛选与分组均值','M1、M2：材料已核'),('继续查个体与采样依据','M3：仍待补资料')]
 expected_pairs={(nodes[a],nodes[b]) for a,b in expected_edges}
 actual_pairs={(c.find('.//a:stCxn',ns).attrib['id'],c.find('.//a:endCxn',ns).attrib['id']) for c in flow.findall('.//p:cxnSp',ns)}
 assert actual_pairs==expected_pairs
 chartfiles=[n for n in z.namelist() if re.fullmatch(r'ppt/(?:slides/)?charts/chart\d+\.xml',n)];assert len(chartfiles)==1
 x=E.fromstring(z.read(chartfiles[0]));vals=[float(t.text) for t in x.findall('.//c:ser/c:val/c:numRef/c:numCache/c:pt/c:v',ns)];assert len(vals)==2 and all(abs(a-b)<1e-9 for a,b in zip(vals,means))
 assert x.find('.//c:valAx/c:scaling/c:min',ns).attrib['val']=='0'
 books=[n for n in z.namelist() if n.startswith('ppt/embeddings/') and n.endswith('.xlsx')];assert len(books)==1
 book=load_workbook(io.BytesIO(z.read(books[0])),data_only=True);nums=[c.value for ws in book for row in ws for c in row if isinstance(c.value,(int,float))]
 assert all(any(abs(v-m)<1e-9 for v in nums) for m in means)
 assert any(z.read(n)==(packet/'word-preview/page-1.png').read_bytes() for n in z.namelist() if n.startswith('ppt/media/') and n.endswith('.png'))
pdf=r/'renders'/rev/(p.stem+'.pdf');assert len(PdfReader(pdf).pages)==10 and len(list(pdf.parent.glob('slide-*.png')))==10
subprocess.run([sys.executable,str(r/'software-packet/check_software.py')],check=True,capture_output=True)
software=load(r/'software-packet/checks.json')
assert all(t in texts[7] for t in ['SW01','344','146','73','尚未确认','未申请','无登记号或证书'])
unchanged=[];same_previews=[]
prior=next((r/'output').glob('*s12-r5.pptx'),None)
if prior:
 with zipfile.ZipFile(prior) as old,zipfile.ZipFile(p) as new:
  for i in ([2,3,6,8,10] if rev=='s12-r8' else [1,2,3,4,6,8,10]):
   name=f'ppt/slides/slide{i}.xml';ot=[t.text for t in E.fromstring(old.read(name)).findall('.//a:t',ns)];nt=[t.text for t in E.fromstring(new.read(name)).findall('.//a:t',ns)];assert ot==nt;unchanged.append(i)
   oldpng=r/'renders/s12-r5'/f'slide-{i:02d}.png'
   if oldpng.exists() and oldpng.read_bytes()==(r/'renders'/rev/f'slide-{i:02d}.png').read_bytes():same_previews.append(i)
vpath=r/'.build'/rev/'validation.json'
v=load(vpath if vpath.exists() else r/'validation.json');digest=sha(p);assert v['finalSha256']==digest and v['packageIntegrity']['findingCount']==v['presentationLayout']['findingCount']==0
(r/'validation.json').write_text(json.dumps(v,ensure_ascii=False,indent=2),'utf-8')
q=dict(revision=rev,pages=10,notes=10,pdf_pages=10,preview_images=10,editable_connector_diagram_pages=diagrams,native_table_pages=tables,native_chart_pages=charts,embedded_workbooks=1,actual_word_preview_pages=images,word_pages=3,attachment_files=9,retained_rows_with_source_lines=146,actual_file_execution=True,source_records=344,analysis_records=146,mean_difference_mm=calc['mean_difference_mm'],verified_material_items=2,pending_items=1,actual_negative_cases_blocked=4,external_model_called=False,scientific_validation='not_performed',pptx_sha256=digest,software_snapshot=software,unchanged_slide_text_vs_s12_r5=unchanged,identical_preview_bytes_vs_s12_r5=same_previews,scope='公开CSV独立复算、材料状态与哈希、本轮4次真实失败执行记录、9项附件文件哈希及原CSV逐行对应、Word内容、PPT原生对象及导出检查；新增软件快照实际执行及台账核对，不证明贡献、权属、登记适格性、科研验收、财务审核或云端模型全系统运行。')
(r/'checks.json').write_text(json.dumps(q,ensure_ascii=False,indent=2),'utf-8');print(json.dumps(q,ensure_ascii=False,indent=2))
