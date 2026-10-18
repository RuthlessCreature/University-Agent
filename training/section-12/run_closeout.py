"""Section 12 teaching evidence review. Actual file checks, no external model.

Reads CSV/JSON only. Does not implement PDF/OCR, official submission, financial
review, scientific sign-off or the deployed Cloudflare Closeout workflow.
"""
from pathlib import Path
import csv,json,hashlib,statistics,datetime,sys
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.oxml.ns import qn

class ReviewError(ValueError):pass
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def dump(p,v):Path(p).write_text(json.dumps(v,ensure_ascii=False,indent=2),'utf-8')
def load(p):return json.loads(Path(p).read_text('utf-8-sig'))

def run(packet,out):
 packet=Path(packet);out=Path(out);out.mkdir(parents=True,exist_ok=True)
 events=[]
 def event(stage,state):events.append(dict(stage=stage,state=state,utc=datetime.datetime.now(datetime.timezone.utc).isoformat()))
 try:
  event('intake','running');task=load(packet/'input/taskbook.json');source=load(packet/'input/source.json');data=packet/'input/penguins.csv'
  if sha(data)!=source['sha256']:raise ReviewError('source_hash_mismatch')
  if task.get('material_type')!='teaching-simulation':raise ReviewError('teaching_scope_required')
  ids=[q['id'] for q in task['requirements']]
  if len(ids)!=len(set(ids)):raise ReviewError('duplicate_requirement')
  if set(ids)!={'M1','M2','M3'}:raise ReviewError('unsupported_requirement_set')
  req={q['id']:q for q in task['requirements']}
  if [req[k]['rule'] for k in ['M1','M2','M3']]!=['source_and_selection','group_means','sampling_evidence']:raise ReviewError('unsupported_rule')
  with data.open(encoding='utf-8-sig',newline='') as f:
   reader=csv.DictReader(f);fields=reader.fieldnames;raw=list(reader)
  adelie=[(i,q) for i,q in enumerate(raw,2) if q['species']=='Adelie']
  kept=[(i,q) for i,q in adelie if q['sex'] in ['female','male'] and q['flipper_length_mm'] not in ['', 'NA']]
  groups={s:[float(q['flipper_length_mm']) for _,q in kept if q['sex']==s] for s in ['female','male']}
  calc=dict(raw_rows=len(raw),adelie_rows=len(adelie),analysis_rows=len(kept),excluded_csv_lines=[i for i,q in adelie if not any(j==i for j,_ in kept)],female_n=len(groups['female']),male_n=len(groups['male']),female_mean_mm=statistics.mean(groups['female']),male_mean_mm=statistics.mean(groups['male']))
  calc['mean_difference_mm']=calc['male_mean_mm']-calc['female_mean_mm'];calc['source_sha256']=sha(data)
  if any(calc[k]!=v for k,v in req['M1']['expected'].items()):raise ReviewError('selection_result_mismatch')
  if any(abs(calc[k]-v)>1e-10 for k,v in req['M2']['expected'].items()):raise ReviewError('mean_result_mismatch')
  missing=[k for k in req['M3']['required_fields'] if k not in fields]
  if not missing:raise ReviewError('sampling_evidence_requires_manual_review')
  matrix=[dict(id='M1',criterion=req['M1']['criterion'],status='VERIFIED',status_zh='材料已核',evidence='input/source.json、input/penguins.csv',locator='CSV列名及全部344条记录',result=f"344条原记录，Adelie 152条，筛选后146条",limit='核对公开来源与筛选，不确认原始采集质量'),dict(id='M2',criterion=req['M2']['criterion'],status='VERIFIED',status_zh='材料已核',evidence='output/calculation.json、input/penguins.csv',locator='sex、flipper_length_mm列；Adelie记录',result=f"雌雄各73条，均值差{calc['mean_difference_mm']:.2f} mm",limit='描述统计，未证明独立性、因果或总体外推'),dict(id='M3',criterion=req['M3']['criterion'],status='PENDING',status_zh='待补资料',evidence='input/penguins.csv',locator='完整列名列表',result='简化CSV缺个体标识及抽样设计资料',limit='需回到原始数据与采样记录核对，当前不判断已完成')]
  if task.get('overall_completed_claim') and any(q['status']!='VERIFIED' for q in matrix):raise ReviewError('unresolved_completion_claim')
  event('source_and_calculation','passed');event('evidence_mapping','pending_sampling_evidence')
  dump(out/'calculation.json',calc);dump(out/'evidence-matrix.json',matrix)
  with (out/'evidence-matrix.csv').open('w',encoding='utf-8-sig',newline='') as f:
   w=csv.DictWriter(f,fieldnames=list(matrix[0]));w.writeheader();w.writerows(matrix)
  issues=[dict(id='SAMPLING',status='PENDING',issue='缺个体标识与抽样设计核对资料'),dict(id='SIGNOFF',status='PENDING',issue='未取得研究者对科学解释与当前版本的核对')]
  dump(out/'issues.json',issues);dump(out/'excluded-claims.json',task['excluded_claims'])
  d=Document();normal=d.styles['Normal'];normal.font.name='Microsoft YaHei';normal.font.size=Pt(10.5);normal.paragraph_format.space_after=Pt(7);normal._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
  for name in ['Title','Heading 1','Heading 2']:
   st=d.styles[name];st.font.name='Microsoft YaHei';st.font.color.rgb=RGBColor(0,0,0);st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
  d.styles['Title'].font.size=Pt(21);d.styles['Heading 1'].font.size=Pt(14)
  sec=d.sections[0];sec.top_margin=sec.bottom_margin=Cm(1.8);sec.left_margin=sec.right_margin=Cm(2)
  d.add_paragraph('企鹅公开数据结题材料核对示例',style='Title')
  d.add_paragraph('教学材料 供讲师演示证据整理')
  d.add_paragraph('本次已核对公开数据来源、筛选记录和分组均值，个体重复记录与抽样依据仍待补充。这里的“材料已核”只对应列明的文件检查，不表示科学结论获批或正式项目通过验收。')
  d.add_paragraph('任务与材料对应',style='Heading 1')
  table=d.add_table(rows=1,cols=3);table.style='Table Grid'
  for c,t in zip(table.rows[0].cells,['条目与核对要求','当前依据与结果','材料状态']):c.text=t
  for q in matrix:
   cells=table.add_row().cells
   loc={'M1':'来源表与原始CSV','M2':'计算结果与原始CSV','M3':'简化CSV的列名'}[q['id']]
   for c,t in zip(cells,[q['id']+' '+q['criterion'],q['result']+'\n依据：'+loc,q['status_zh']]):c.text=t
  d.add_paragraph('实际复算结果',style='Heading 1')
  d.add_paragraph(f"公开简化数据共{calc['raw_rows']}条，Adelie共{calc['adelie_rows']}条。保留性别与鳍长均有记录的{calc['analysis_rows']}条，其中雌雄各73条。雌性平均鳍长{calc['female_mean_mm']:.2f} mm，雄性{calc['male_mean_mm']:.2f} mm，雄性均值高{calc['mean_difference_mm']:.2f} mm。上述结果描述本次所选记录，不能直接解释因果机制。")
  d.add_paragraph('数据来自Gorman及Palmer Station LTER公开项目，经palmerpenguins提供简化版本。本课程只作二次计算，未开展野外采集，也没有已发表论文或基金项目资助成果的主张。')
  d.add_page_break()
  d.add_paragraph('缺口及后续核对',style='Heading 1')
  d.add_paragraph('简化CSV不含individual_id和sampling_design字段。需要进一步查看原始个体标识与采样记录，检查同一个体重复观测、分组和独立性假设。当前不把这一项写为完成，也不以现有均值差支持因果或总体外推。')
  d.add_paragraph('本案例任务书为课程自编材料核对要求，借已有数据演示结题组织，未作为事先登记的研究方案。没有经费账目、拨款记录、论文发表、专利授权或正式审核记录，相关内容不填写。')
  d.add_paragraph('可据此使用的描述',style='Heading 1')
  d.add_paragraph('已对公开数据中的Adelie记录完成筛选和均值复算，保存了原始CSV、程序及结果。个体标识和采样依赖尚未核对，结论限于所选记录的描述。')
  d.add_paragraph('应排除的说法',style='Heading 1')
  for q in task['excluded_claims']:d.add_paragraph(q['claim']+'。该说法没有支持：'+q['reason']+'。')
  d.add_paragraph('来源与附件定位',style='Heading 1')
  for t in ['公开数据说明 https://allisonhorst.github.io/palmerpenguins/','数据与来源 input/penguins.csv 和 input/source.json','自编核对要求 input/taskbook.json 的M1至M3','实际计算 output/calculation.json；证据表 output/evidence-matrix.csv','执行器 run_closeout.py；文件版本 output/run.json'] :d.add_paragraph(t)
  d.add_paragraph('本文件为AI编写的教学材料核对示例。国家自然科学基金委2026年通告第五部分第（二）项对生成式人工智能直接生成的项目进展和结题报告作出限制。本文件不能作为该类正式报告提交。')
  for tree in [d._element,d.styles.element]:
   for el in list(tree.iter(qn('w:pBdr'))):el.getparent().remove(el)
  artifact=out/'materials-review.docx';d.save(artifact);event('docx','written')
  files=[packet/'input/taskbook.json',packet/'input/source.json',data]
  result=dict(status='materials_review_required',material_type='teaching-simulation',external_model_called=False,human_approved=False,scientific_validation='not_performed',official_submission=False,input_hashes={p.name:sha(p) for p in files},executor_sha256=sha(__file__),artifact=artifact.name,artifact_sha256=sha(artifact),verified_material_items=2,pending_items=1,events=events)
  dump(out/'run.json',result);return result
 except Exception as e:
  dump(out/'run.json',dict(status='failed',error=str(e),events=events));raise

if __name__=='__main__':
 r=Path(__file__).resolve().parent
 result=run(r/'closeout-packet',r/'closeout-packet/output')
 dump(r/'execution-evidence.json',result);print(result['status'])
