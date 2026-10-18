from pathlib import Path
from xml.etree import ElementTree as E
import csv,json,hashlib,zipfile,re,statistics,io
from docx import Document
from pypdf import PdfReader
from openpyxl import load_workbook
r=Path(__file__).resolve().parent;rev='s13-r6';packet=r/'communication-packet';p=next((r/'output').glob('*'+rev+'.pptx'))
def load(p):return json.loads(p.read_text('utf-8'))
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
run=load(packet/'output/run.json');assert run==load(r/'execution-evidence.json')
assert all(sha(packet/'input'/name)==h for name,h in run['input_hashes'].items())
assert sha(r/'make_packet.py')==run['executor_sha256'] and sha(packet/'output/汇报与传播文字示例.docx')==run['artifact_sha256']
assert run['status']=='teaching_draft_for_review' and not run['external_model_called'] and not run['researcher_approved'] and not run['published'] and not run['platform_account_accessed']
with (packet/'input/penguins.csv').open(encoding='utf-8-sig',newline='') as f:raw=list(csv.DictReader(f))
clean=[q for q in raw if q['species']=='Adelie' and q['sex'] in ['female','male'] and q['flipper_length_mm'] not in ['','NA']]
assert len(raw)==344 and len(clean)==146
groups={s:[float(q['flipper_length_mm']) for q in clean if q['sex']==s] for s in ['female','male']};means=[statistics.fmean(groups[s]) for s in ['female','male']]
facts=load(packet/'output/facts.json');assert all(abs(means[i]-facts[s+'_mean_mm'])<1e-10 for i,s in enumerate(['female','male'])) and abs(means[1]-means[0]-facts['difference_mm'])<1e-10
edges=facts['histogram']['edges'];h={s:[sum(lo<=v<hi or (i==len(edges)-2 and v==hi) for v in groups[s]) for i,(lo,hi) in enumerate(zip(edges,edges[1:]))] for s in groups};assert h==facts['histogram']['counts'] and sum(h['female'])==sum(h['male'])==73
brief=load(packet/'input/brief.json');assert len(brief['versions'])==3
assert '平均比雌性长4.62 mm' in brief['versions'][2]['text']
doc=Document(packet/'output/汇报与传播文字示例.docx');body=''.join(q.text.replace('\n','') for q in doc.paragraphs).replace('\u00a0',' ')
for q in brief['versions']:assert q['audience'] in body and q['text'] in body and all(s in q['text'] for s in ['146','73','4.62 mm'])
assert brief['humanizer_example']['after'] in body and 'AI编写' in body and '{{' not in body
assert len(PdfReader(packet/'word-preview/汇报与传播文字示例.pdf').pages)==2 and len(list((packet/'word-preview').glob('page-*.png')))==2
lesson=load(r/'lesson.json');ns={'a':'http://schemas.openxmlformats.org/drawingml/2006/main','p':'http://schemas.openxmlformats.org/presentationml/2006/main','c':'http://schemas.openxmlformats.org/drawingml/2006/chart'};diagrams=[];tables=[];charts=[];images=[];texts=[]
with zipfile.ZipFile(p) as z:
 slides=[n for n in z.namelist() if re.fullmatch(r'ppt/slides/slide\d+\.xml',n)];notes=[n for n in z.namelist() if re.fullmatch(r'ppt/notesSlides/notesSlide\d+\.xml',n)];assert len(slides)==len(notes)==len(lesson)==10
 for i in range(1,11):
  x=E.fromstring(z.read(f'ppt/slides/slide{i}.xml'));text=' '.join(t.text or '' for t in x.findall('.//a:t',ns)).replace('\u00a0',' ');texts.append(text);assert lesson[i-1]['title'] in text
  assert all(t not in text for t in ['TODO','课堂练习','占位'])
  if len(x.findall('.//p:cxnSp',ns))>=2:diagrams.append(i)
  if x.findall('.//a:tbl',ns):tables.append(i)
  if x.findall('.//c:chart',ns):charts.append(i)
  if x.findall('.//p:pic',ns):images.append(i)
  nt=z.read(f'ppt/notesSlides/notesSlide{i}.xml').decode();assert 'http' in nt or 'local:' in nt
 assert tables==[5] and charts==[4] and images==[3,9]
 for q in brief['versions']:assert q['text'] in texts[4]
 chartfiles=[n for n in z.namelist() if re.fullmatch(r'ppt/(?:slides/)?charts/chart\d+\.xml',n)];assert len(chartfiles)==1
 x=E.fromstring(z.read(chartfiles[0]));vals=[[float(t.text) for t in ser.findall('.//c:val/c:numRef/c:numCache/c:pt/c:v',ns)] for ser in x.findall('.//c:ser',ns)];assert vals==[h['female'],h['male']]
 assert x.find('.//c:valAx/c:scaling/c:min',ns).attrib['val']=='0'
 books=[n for n in z.namelist() if n.startswith('ppt/embeddings/') and n.endswith('.xlsx')];assert len(books)==1
 book=load_workbook(io.BytesIO(z.read(books[0])),data_only=True);rows=list(book.active.values);assert len(rows)>=10
 for expected in [h['female'],h['male']]:assert any([row[col] for row in rows[1:10]]==expected for col in range(len(rows[0])))
 for file in [packet/'word-preview/page-1.png',r/'assets/uf-2025-page-23.png']:
  assert any(z.read(n)==file.read_bytes() for n in z.namelist() if n.startswith('ppt/media/') and n.endswith('.png'))
pdf=r/'renders'/rev/(p.stem+'.pdf');assert len(PdfReader(pdf).pages)==10 and len(list(pdf.parent.glob('slide-*.png')))==10
v=load(r/'.build'/rev/'validation.json');digest=sha(p);assert v['finalSha256']==digest and v['packageIntegrity']['findingCount']==v['presentationLayout']['findingCount']==0
(r/'validation.json').write_text(json.dumps(v,ensure_ascii=False,indent=2),'utf-8')
q=dict(revision=rev,pages=10,notes=10,pdf_pages=10,preview_images=10,editable_connector_diagram_pages=diagrams,native_table_pages=tables,native_chart_pages=charts,embedded_workbooks=1,sourced_figure_page=3,actual_word_preview_page=9,word_pages=2,reader_versions=3,source_records=344,analysis_records=146,histogram_records_per_group=73,actual_file_execution=True,external_model_called=False,platform_account_accessed=False,researcher_approved=False,pptx_sha256=digest,scope='原CSV分组与分箱独立复算、三个完整段落及比较对象、Word与源图字节、原生图表表格及导出检查；数字匹配不等于语义审查、科学验收或实际平台发布。')
(r/'checks.json').write_text(json.dumps(q,ensure_ascii=False,indent=2),'utf-8');print(json.dumps(q,ensure_ascii=False,indent=2))
