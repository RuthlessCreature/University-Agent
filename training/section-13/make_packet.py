"""Recompute source facts and assemble authored teaching text; no model call."""
from pathlib import Path
import csv,json,hashlib,statistics,datetime
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
r=Path(__file__).resolve().parent;p=r/'communication-packet';out=p/'output';out.mkdir(parents=True,exist_ok=True)
def sha(f):return hashlib.sha256(f.read_bytes()).hexdigest()
def load(f):return json.loads(f.read_text('utf-8-sig'))
def dump(f,v):f.write_text(json.dumps(v,ensure_ascii=False,indent=2),'utf-8')
src=load(p/'input/source.json');brief=load(p/'input/brief.json');assert sha(p/'input/penguins.csv')==src['sha256']
with (p/'input/penguins.csv').open(encoding='utf-8-sig',newline='') as f:raw=list(csv.DictReader(f))
clean=[q for q in raw if q['species']=='Adelie' and q['sex'] in ['female','male'] and q['flipper_length_mm'] not in ['','NA']]
groups={s:[float(q['flipper_length_mm']) for q in clean if q['sex']==s] for s in ['female','male']}
edges=list(range(170,216,5));hist={s:[sum(lo<=v<hi or (i==len(edges)-2 and v==hi) for v in groups[s]) for i,(lo,hi) in enumerate(zip(edges,edges[1:]))] for s in groups}
facts=dict(raw_n=len(raw),analysis_n=len(clean),female_n=len(groups['female']),male_n=len(groups['male']),female_mean_mm=statistics.mean(groups['female']),male_mean_mm=statistics.mean(groups['male']),histogram=dict(edges=edges,counts=hist))
facts['difference_mm']=facts['male_mean_mm']-facts['female_mean_mm'];assert sum(hist['female'])==sum(hist['male'])==73
for k,v in brief['fact_contract'].items():assert (round(facts[k],2) if k=='difference_mm' else facts[k])==v
for version in brief['versions']:
 assert all(token in version['text'] for token in ['146','73','4.62 mm'])
assert brief['release_state']=='not_published' and brief['researcher_approved']==False
dump(out/'facts.json',facts)
audit=dict(scope='锁定字段和数字匹配；语义、来源归属与限定另经逐句阅读，不是自动科学审查',fact_contract=brief['fact_contract'],versions=[dict(audience=q['audience'],numeric_tokens_preserved=True) for q in brief['versions']],release_state='not_published',researcher_approved=False)
dump(out/'expression-check.json',audit)
md='# 企鹅公开数据复算的汇报与传播示例\n\nAI编写的教学表达，未发布，未取得研究者确认。原始采集属于公开项目，本课程只作二次计算。\n\n'
for q in brief['versions']:md+='## '+q['audience']+'\n\n'+q['text']+'\n\n'
md+='## 表达修订\n\n教学反例：'+brief['humanizer_example']['before']+'\n\n修订后：'+brief['humanizer_example']['after']+'\n\n'
md+='## 范围与出处\n\n'+'\n\n'.join(brief['limits'])+'\n\n公开数据：https://allisonhorst.github.io/palmerpenguins/\n\n输入：input/penguins.csv；计算：output/facts.json；文本：input/brief.json。\n'
(out/'传播文字示例.md').write_text(md,'utf-8')
d=Document();normal=d.styles['Normal'];normal.font.name='Microsoft YaHei';normal.font.size=Pt(10.5);normal.paragraph_format.space_after=Pt(8);normal._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
for sn in ['Title','Heading 1','Heading 2']:
 st=d.styles[sn];st.font.name='Microsoft YaHei';st.font.color.rgb=RGBColor(0,0,0);st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
d.styles['Title'].font.size=Pt(21);d.styles['Heading 1'].font.size=Pt(14)
s=d.sections[0];s.top_margin=s.bottom_margin=Cm(1.8);s.left_margin=s.right_margin=Cm(2)
d.add_paragraph('企鹅公开数据复算的汇报与传播示例',style='Title')
d.add_paragraph('教学稿 供表达比较')
d.add_paragraph('三个版本使用同一组公开数据复算结果，分别面向同行、评审人员和普通读者。数据由原研究项目采集，本课程只作二次计算。文本为AI编写的教学表达，尚未取得研究者确认。')
d.add_paragraph('各版本共用的事实',style='Heading 1')
t=d.add_table(rows=1,cols=2);t.style='Table Grid'
for c,v in zip(t.rows[0].cells,['核对内容','实际结果']):c.text=v
for a,b in [('所选记录','146条Adelie，雌雄各73条'),('平均鳍长','雌性187.79 mm，雄性192.41 mm'),('均值差与方向','雄性平均比雌性高4.62 mm'),('解释范围','描述所选记录，未核对采样依赖或识别因果')]:
 for c,v in zip(t.add_row().cells,[a,b]):c.text=v
for q in brief['versions'][:2]:d.add_paragraph(q['audience'],style='Heading 1');d.add_paragraph(q['text'])
q=brief['versions'][2];d.add_paragraph(q['audience'],style='Heading 1');d.add_paragraph(q['text']);d.add_page_break()
d.add_paragraph('表达修订中的事实核对',style='Heading 1')
d.add_paragraph('教学反例 '+brief['humanizer_example']['before'])
d.add_paragraph('修订后 '+brief['humanizer_example']['after'])
d.add_paragraph('修改删去无证据的突破评价，将因果机制改回描述统计，并补回数据归属、分析对象和单位。语言更顺之后，数字、单位、比较方向和限定仍需逐项核对。')
d.add_paragraph('发布前需要核对的内容',style='Heading 1')
d.add_paragraph('核对署名与成果归属、文件版本和公开权限，确认能否分享全文、数据或图片。引用他人的成果应回到原始来源，并按许可及平台规则选择可公开版本。本例未登录科研平台，也未代签或发布内容。')
d.add_paragraph('来源与文件定位',style='Heading 1')
for v in ['公开数据 https://allisonhorst.github.io/palmerpenguins/','来源及原始CSV input/source.json 和 input/penguins.csv','复算结果 output/facts.json；各版本文字 input/brief.json','ORCID https://info.orcid.org/researchers/','ResearchGate版权说明 https://help.researchgate.net/copyright/copyright-and-researchgate']:d.add_paragraph(v)
for paragraph in list(d.paragraphs)+[p for t in d.tables for row in t.rows for c in row.cells for p in c.paragraphs]:
 wrap=OxmlElement('w:wordWrap');wrap.set(qn('w:val'),'0');paragraph._p.get_or_add_pPr().append(wrap)
 for run in paragraph.runs:
  if '4.62 mm' in run.text:run.text=run.text.replace('4.62 mm','\n4.62 mm')
for tree in [d._element,d.styles.element]:
 for el in list(tree.iter(qn('w:pBdr'))):el.getparent().remove(el)
artifact=out/'汇报与传播文字示例.docx';d.save(artifact)
record=dict(status='teaching_draft_for_review',utc=datetime.datetime.now(datetime.timezone.utc).isoformat(),input_hashes={f.name:sha(f) for f in [p/'input/penguins.csv',p/'input/source.json',p/'input/brief.json']},executor_sha256=sha(Path(__file__)),artifact_sha256=sha(artifact),external_model_called=False,researcher_approved=False,published=False,platform_account_accessed=False,scope='公开CSV复算、锁定数字匹配与现成教学文字整理，不是自主传播Agent运行或正式发表')
dump(out/'run.json',record);dump(r/'execution-evidence.json',record);print(record['status'])
