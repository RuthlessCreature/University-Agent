from pathlib import Path
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml.ns import qn
from PIL import Image
R=Path(__file__).resolve().parent;O=R/'output';O.mkdir(exist_ok=True)
d=Document();s=d.sections[0];s.top_margin=s.bottom_margin=Cm(2);s.left_margin=s.right_margin=Cm(2.2);s.page_height=Cm(29.7);s.page_width=Cm(21)
for name in ['Normal','Title','Heading 1']:
    st=d.styles[name];st.font.name='Microsoft YaHei';st._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei');st.font.color.rgb=RGBColor(0,0,0)
d.styles['Title'].font.size=Pt(21);d.styles['Normal'].font.size=Pt(11)
d.styles['Normal'].paragraph_format.line_spacing=1.3;d.styles['Normal'].paragraph_format.space_after=Pt(8)
d.styles['Heading 1'].font.size=Pt(13)
d.add_paragraph('一篇论文的阅读笔记','Title')
d.add_paragraph('课堂示例   2026 年 10 月 17 日')
d.add_paragraph('这份笔记根据公开论文整理，含 AI 辅助编写。它回答一个具体问题，不代替对整篇论文和相关研究的阅读。')
items=[('我问的问题','AI 帮老师生成作业评分和评语，能让学生学习成绩提高吗？'),('这篇论文能告诉我什么','这篇论文讨论 AI 辅助评分和反馈，没有直接检查学生是否学得更好。因此，不能把这篇论文当作“AI 评语已经提高学生成绩”的证据。'),('我是在哪里核对的','原 PDF 第 2 页，Introduction（引言）。在文件中搜索 learning outcomes，可以找到相关说明。关键短语为 does not directly assess student learning outcomes，意思是没有直接评价学生的学习结果。'),('论文出处','Heinrich 等，2025。AI-assisted grading and personalized feedback in large political science classes: Results from randomized controlled trials. PLOS ONE 20(8): e0328041。\n固定编号：https://doi.org/10.1371/journal.pone.0328041'),('接下来还需要找什么','如果要判断是否提高成绩，还需要查找实际测量学生学习结果的研究。没有测这一项，不能解释为“已经证明有效”或“已经证明无效”。')]
for title,body in items:d.add_paragraph(title,'Heading 1');d.add_paragraph(body)
for parent in [d.styles.element,d.element]:
    for el in list(parent.iter(qn('w:pBdr'))):el.getparent().remove(el)
d.save(O/'论文阅读笔记.docx')
im=Image.open(R/'assets/page1.png');im.crop((350,168,1045,384)).save(R/'assets/paper-title.png')
print('Word note and disclosed title crop saved')
