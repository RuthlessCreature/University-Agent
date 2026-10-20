"""Export teaching revision files from the existing Markdown and review matrix.

The marked copy uses visible run formatting, not Word Track Changes. The manifest
records both accepted and deleted paragraph projections for an independent check.
"""
from pathlib import Path
import csv, difflib, hashlib, json, re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
SRC = ROOT / 'writing-packet'
OUT = SRC / 'word-r1'
OUT.mkdir(exist_ok=True)

def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()

def document(title, intro):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2)
    sec.left_margin = sec.right_margin = Cm(2.1)
    for name, size in [('Normal', 11), ('Title', 18), ('Heading 1', 13), ('Heading 2', 12)]:
        st = doc.styles[name]
        st.font.name = 'Microsoft YaHei'
        st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.space_before = Pt(6 if name.startswith('Heading') else 0)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.line_spacing = 1.1
    for border in list(doc.styles.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    doc.add_paragraph(title, 'Title')
    doc.add_paragraph(intro)
    return doc

def line(doc, text, status=None):
    text = text.strip()
    if not text:
        return
    level = len(text) - len(text.lstrip('#'))
    plain = text.lstrip('#').strip() if level else text
    p = doc.add_paragraph(style='Heading 1' if level else 'Normal')
    run = p.add_run(plain)
    if status == 'delete':
        run.font.color.rgb = RGBColor.from_string('963F3C')
        run.font.strike = True
    elif status == 'insert':
        run.font.color.rgb = RGBColor.from_string('246A52')
        run.font.underline = True
    return p

def plain_lines(text):
    return [s for s in text.splitlines() if s.strip()]

r0 = (SRC/'manuscript-r0.md').read_text('utf-8')
r1 = (SRC/'manuscript-r1.md').read_text('utf-8')
old, new = plain_lines(r0), plain_lines(r1)
clean = document('修订稿清稿 教学示例', '公开企鹅数据的二次分析写作示例。本文未投稿，不主张新发现。方括号中的段落编号用于与模拟审稿回复核对。')
for s in new:
    line(clean, s)
clean.save(OUT/'修订稿_清稿.docx')

marked = document('修订稿改动标记 教学示例', '比较对象为随包的 manuscript-r0.md 与 manuscript-r1.md。红色删除线表示原稿删除，绿色下划线表示修订新增，黑色表示相同文字。这里使用可编辑的字体标记，没有启用 Word 修订模式。原稿含故意设置的错误结论，仅供辨析。')
projection_old, projection_new = [], []
for tag, i, j, a, b in difflib.SequenceMatcher(a=old, b=new, autojunk=False).get_opcodes():
    if tag == 'equal':
        for s in old[i:j]:
            line(marked, s)
            projection_old.append(s)
            projection_new.append(s)
    else:
        for s in old[i:j]:
            line(marked, s, 'delete')
            projection_old.append(s)
        for s in new[a:b]:
            line(marked, s, 'insert')
            projection_new.append(s)
assert projection_old == old and projection_new == new
marked.save(OUT/'修订稿_改动标记.docx')

matrix = list(csv.DictReader((SRC/'response-matrix.csv').open(encoding='utf-8-sig')))
reply = document('逐条审稿回复 教学模拟', '本文件对应课程自编的三条模拟审稿意见。两项正文修改已写入教学稿，新增采样及巢内相关性检验尚未完成。所有定位均使用清稿中的段落编号，未经历真实编辑或审稿流程。')
for q in matrix:
    reply.add_paragraph(q['id'] + '  ' + ('补充方法说明' if q['id']=='R1.1' else '收回因果结论' if q['id']=='R1.2' else '尚未完成的补充研究'), 'Heading 1')
    for label, text in [('模拟意见', q['comment']), ('回复', q['response']), ('修改位置', q['anchor']+'，见修订稿_清稿.docx及修订稿_改动标记.docx。'), ('状态', q['status']), ('核对材料', q['evidence'])]:
        p=reply.add_paragraph()
        p.add_run(label+'：').bold=True
        p.add_run(text)
reply.save(OUT/'逐条审稿回复.docx')

manifest = {
    'revision': 'word-r1',
    'material_status': '公开数据二次分析教学稿；审稿意见为模拟；未投稿',
    'marking': 'red strike deletions and green underline insertions; not Word Track Changes',
    'source_hash_method': 'SHA256 of UTF-8 text with LF newlines; independent of checkout CRLF conversion',
    'sources': {p.name: hashlib.sha256(p.read_text('utf-8').encode('utf-8')).hexdigest() for p in [SRC/'manuscript-r0.md', SRC/'manuscript-r1.md', SRC/'response-matrix.csv']},
    'exports': {p.name: sha(p) for p in sorted(OUT.glob('*.docx'))},
    'accepted_plain_lines': [s.lstrip('#').strip() for s in new],
    'rejected_plain_lines': [s.lstrip('#').strip() for s in old],
    'reviews': [{'id':q['id'], 'anchors':re.findall(r'\[[A-Z][0-9]+\]',q['anchor']), 'status':q['status'], 'evidence':q['evidence']} for q in matrix],
    'scope': '文件生成与版本对应；不代表研究验收、全文投稿合规或期刊接受'
}
(OUT/'export-manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), 'utf-8')
print(json.dumps({'exports': manifest['exports'], 'reviews': len(matrix)}, ensure_ascii=False, indent=2))
